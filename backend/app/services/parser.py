from io import BytesIO
import os

from pypdf import PdfReader
from docx import Document


def extract_pdf_text(file_content: bytes) -> str:
    pdf_file = BytesIO(file_content)
    reader = PdfReader(pdf_file)

    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text).strip()


def extract_docx_text(file_content: bytes) -> str:
    docx_file = BytesIO(file_content)
    document = Document(docx_file)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    # Also pull text from tables, which many resumes use.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    text.append(cell_text)

    return "\n".join(text).strip()


def extract_txt_text(file_content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue

    raise ValueError("Could not decode text resume. Please save it as UTF-8.")


def _normalize_extension(filename_or_extension: str) -> str:
    value = (filename_or_extension or "").strip().lower()

    if not value:
        return ""

    # Accept either a full filename ("resume.pdf") or an extension (".pdf" / "pdf").
    if "." in value and not value.startswith("."):
        value = os.path.splitext(value)[1]

    if value and not value.startswith("."):
        value = f".{value}"

    return value


def extract_text(file_content: bytes, filename_or_extension: str) -> str:
    extension = _normalize_extension(filename_or_extension)

    if extension == ".pdf":
        text = extract_pdf_text(file_content)
    elif extension == ".docx":
        text = extract_docx_text(file_content)
    elif extension == ".txt":
        text = extract_txt_text(file_content)
    elif extension == ".doc":
        raise ValueError(
            "Old .doc files are not supported. Please save your resume as PDF or DOCX."
        )
    else:
        raise ValueError(
            "Unsupported file format. Only PDF, DOCX, and TXT are supported."
        )

    if not text:
        raise ValueError(
            "Could not extract any text from the resume. "
            "Try a text-based PDF or DOCX file."
        )

    return text
